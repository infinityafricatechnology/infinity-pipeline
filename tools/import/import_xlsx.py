# /// script
# requires-python = ">=3.11"
# dependencies = ["openpyxl==3.1.5"]
# ///
"""Convertit le classeur PIPELINE_IAT_V3.xlsx en données de départ (JSON) pour l'application.

Usage : uv run tools/import/import_xlsx.py docs/PIPELINE_IAT_V3.xlsx data/seed.json [factures.xlsx]

Le classeur facultatif des renouvellements (un onglet par facture client) complète les factures
(n° FNE, détail des lignes, remise, acompte) et les coordonnées de facturation des comptes.

Les valeurs calculées du classeur (montant pondéré, jours restants, alertes, totaux) ne sont pas
reprises : l'application les recalcule. Seules les saisies le sont.
"""

import json
import re
import sys
import unicodedata
import warnings
from datetime import date, datetime

import openpyxl

warnings.filterwarnings("ignore", category=UserWarning)

ACCOUNT_SHEETS = ["CNAM", "CNPS", "FER", "NSIA", "Orange Business", "NEURONES"]

STAGES = {
    "M0 – À qualifier": "M0",
    "M1 – Besoin identifié": "M1",
    "M2 – Cadrage / découverte": "M2",
    "M3 – Offre en préparation": "M3",
    "M4 – Offre remise": "M4",
    "M5 – Négociation / signature": "M5",
    "M6 – Gagné": "M6",
    "Perdu": "LOST",
    "Clôturé / en cessation": "CLOSED",
}

ACCOUNT_KINDS = {
    "Client existant": "client",
    "Prospect": "prospect",
    "Appel d’offres": "tender",
    "Partenaire / apporteur": "partner",
}

FICHE_FIELDS = {
    "Secteur": "sector",
    "Direction générale": "leadership_summary",
    "Interlocuteur direction": "contact_direction",
    "Interlocuteur clé": "contact_direction",
    "Interlocuteur opérationnel": "contact_operational",
    "Autres interlocuteurs": "contact_operational",
    "Achats / marchés": "contact_purchasing",
    "Responsable IAT": "iat_lead",
    "Relation haut niveau": "executive_sponsor",
    "État de santé": "health",
    "Taux de satisfaction": "satisfaction",
    "Dernier contact": "last_contact",
    "Prochain point prévu": "next_meeting",
    "Notes": "notes",
}


def clean(value):
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, str):
        value = value.strip()
        if value in ("", "—", "-", "N/A"):
            return None
        return value
    return value


def number(value):
    value = clean(value)
    if isinstance(value, (int, float)):
        return value
    return None


def iso_date(value):
    value = clean(value)
    if isinstance(value, str) and re.fullmatch(r"\d{4}-\d{2}-\d{2}", value):
        return value
    return None


def text(value):
    value = clean(value)
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value)


def priority(value):
    value = text(value)
    return value if value in ("P1", "P2", "P3") else None


def slug(name):
    name = re.sub(r"\(.*?\)", "", name).strip().lower()
    name = re.sub(r"[’'«»]", "", name)
    name = unicodedata.normalize("NFKD", name).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "-", name).strip("-")


def rows(ws, first_row, first_col, last_col):
    """Valeurs de la ligne ``first_row`` jusqu'à la première ligne vide (colonnes B=2…)."""
    out = []
    for row in ws.iter_rows(min_row=first_row, min_col=first_col, max_col=last_col):
        values = [cell.value for cell in row]
        if all(clean(v) is None for v in values):
            break
        out.append(values)
    return out


def find_row(ws, predicate, start=1):
    for row in ws.iter_rows(min_row=start, min_col=2, max_col=2):
        cell = row[0]
        if isinstance(cell.value, str) and predicate(cell.value.strip()):
            return cell.row
    return None


class Accounts:
    def __init__(self):
        self.by_key = {}

    def add(self, name, kind, **fields):
        key = slug(name)
        account = self.by_key.get(key)
        if account is None:
            account = {"key": key, "name": name, "kind": kind}
            self.by_key[key] = account
        for field, value in fields.items():
            if value is not None and account.get(field) is None:
                account[field] = value
        return key

    def find(self, label):
        """Le compte dont le nom apparaît dans ``label`` (ex. « Client — fiche CNAM »)."""
        if not label:
            return None
        norm = slug(label)
        for key in sorted(self.by_key, key=len, reverse=True):
            if key and re.search(rf"(^|-){re.escape(key)}(-|$)", norm):
                return key
        return None


def read_account_sheet(ws, accounts):
    name = ws.title
    kind = "partner" if name == "NEURONES" else "client"
    fiche = {}
    start = find_row(ws, lambda v: v.startswith("1. FICHE COMPTE"))
    for label, value, *_ in rows(ws, start + 1, 2, 3):
        label = text(label) or ""
        for prefix, field in FICHE_FIELDS.items():
            if label.startswith(prefix) and fiche.get(field) is None:
                fiche[field] = number(value) if field == "satisfaction" else text(value)
    full_names = {
        "CNAM": "CNAM",
        "CNPS": "CNPS",
        "FER": "FER (Fonds d’entretien routier)",
        "NSIA": "NSIA (Assurances / Assurances Vie)",
        "Orange Business": "Orange Business",
        "NEURONES": "NEURONES",
    }
    key = accounts.add(name, kind, display_name=full_names[name], has_fiche=True, **fiche)

    invoices = []
    start = find_row(ws, lambda v: v.startswith("2. "))
    for r in rows(ws, start + 2, 2, 14):
        if text(r[0]) and text(r[0]).startswith("TOTAL"):
            break
        invoices.append(
            {
                "account": key,
                "product": text(r[0]),
                "vendor": text(r[1]),
                "recurrence": text(r[2]),
                "fiscal_year": number(r[3]) if isinstance(number(r[3]), (int, float)) else None,
                "invoice_date": iso_date(r[4]),
                "amount": number(r[5]),
                "cost_of_sales": number(r[6]),
                "commission": number(r[7]),
                "collected": number(r[10]),
                "comments": text(r[12]) or (text(r[3]) if number(r[3]) is None else None),
            }
        )

    offers = []
    start = find_row(ws, lambda v: v.startswith("3. "))
    for r in rows(ws, start + 2, 2, 14):
        if text(r[0]) and text(r[0]).startswith("TOTAL"):
            break
        offers.append({"title": text(r[0]), "vendor": text(r[1]), "stage": STAGES.get(text(r[2])), "sent_on": iso_date(r[3])})

    opportunities = []
    start = find_row(ws, lambda v: v.startswith("4. "))
    for r in rows(ws, start + 2, 2, 14):
        opportunities.append(
            {
                "account": key,
                "title": text(r[0]),
                "solution": text(r[1]),
                "priority": priority(r[2]),
                "due_date": iso_date(r[3]),
                "owner": text(r[4]),
                "action": text(r[5]),
                "comments": text(r[12]),
            }
        )
    return key, invoices, offers, opportunities


MONTHS = ["janvier", "février", "mars", "avril", "mai", "juin", "juillet", "août", "septembre", "octobre", "novembre", "décembre"]
INVOICE_SHEETS = {"Tableau de bord", "Renouvellements", "Pipeline", "Listes"}


def amount_text(value):
    """« 4 560 000 CFA » → 4560000."""
    if isinstance(value, (int, float)):
        return value
    digits = re.sub(r"[^\d-]", "", value or "")
    return int(digits) if re.fullmatch(r"-?\d+", digits) else None


def money(value):
    return f"{round(value):,}".replace(",", " ")


def read_invoice_sheet(ws):
    """Un onglet facture : date, n°, client facturé (raison sociale, NCC, e-mail…), lignes, remise, acompte."""
    cells = [[clean(v) for v in row] for row in ws.iter_rows(values_only=True)]
    strings = [v for row in cells for v in row if isinstance(v, str)]
    out = {"number": None, "date": None, "client": [], "lines": [], "discount": None, "deposit": None}

    for v in strings:
        if m := re.search(r"FACTURE DE VENTE N° (\S+)", v):
            out["number"] = m.group(1)
        if m := re.search(r"Date et heure : (\d\d)/(\d\d)/(\d{4})", v):
            out["date"] = f"{m.group(3)}-{m.group(2)}-{m.group(1)}"
        if m := re.search(r"Abidjan le (\d+) (\w+) (\d{4})", v):
            out["date"] = f"{m.group(3)}-{MONTHS.index(m.group(2).lower()) + 1:02d}-{int(m.group(1)):02d}"

    header = next(i for i, row in enumerate(cells) if any(isinstance(v, str) and slug(v) == "designation" for v in row))
    head = cells[header]
    client = next((i for row in cells[:header] for i, v in enumerate(row) if v == "CLIENT"), None)
    if client is not None:
        # Facture normalisée (FNE) : colonne « CLIENT » (la colonne « ÉMETTEUR » est IAT).
        for row in cells[:header]:
            if isinstance(row[client], str) and (m := re.match(r"(Nom|NCC|Adresse) : (.+)", row[client])):
                out["client"].append(("NCC " if m.group(1) == "NCC" else "") + m.group(2).strip())
    else:
        # Facture courrier : bloc destinataire entre la date et l'objet.
        for row in cells[:header]:
            for v in row:
                if isinstance(v, str) and not v.startswith(("Objet", "Abidjan le")):
                    v = re.sub(r"^A Monsieur ", "À l'attention de M. ", v)
                    out["client"].append(re.sub(r"^(NCC|Tél|Cél)\s*:\s*", r"\1 ", v))

    designation = next(i for i, v in enumerate(head) if isinstance(v, str) and slug(v) == "designation")
    last = max(i for i, v in enumerate(head) if v is not None)
    unit = next((i for i, v in enumerate(head) if isinstance(v, str) and slug(v).startswith("p-u")), None)
    for row in cells[header + 1 :]:
        if any(isinstance(v, str) and re.match(r"(TOTAL( GENERAL)? HT|Montant Total HT)\s*$", v.strip()) for v in row):
            break
        label = row[designation] if isinstance(row[designation], str) else row[0]
        value = amount_text(row[last]) if not (isinstance(row[last], str) and row[last].startswith("=")) else None
        if value is None and unit is not None and isinstance(row[unit], (int, float)):
            value = row[unit] * (row[unit + 1] or 1) * (1 - (row[unit + 3] or 0))
        if not isinstance(label, str) or label.lower().startswith(("total", "tva")):
            continue
        label = re.sub(r"\s+", " ", label).strip(" ,")
        product = row[0] if designation > 0 and isinstance(row[0], str) and " " in row[0] else None
        if product:
            out["lines"].append((product, None))
        out["lines"].append((label, value) if value else (label, None))

    for row in cells:
        for i, v in enumerate(row):
            if isinstance(v, str) and v.startswith("Taux de remise") and isinstance(row[i + 1], float):
                out["discount"] = row[i + 1]
            if isinstance(v, str) and (m := re.match(r"Acompte reçue? le (\S+)", v)):
                deposit = next(amount_text(x) for x in row[i + 1 :] if x is not None)
                out["deposit"] = (abs(deposit), m.group(1))
    return out


def enrich_from_invoices(path, accounts, invoices):
    """Rapproche chaque onglet facture d'une facture existante (compte + date) et la complète."""
    wb = openpyxl.load_workbook(path, data_only=False)
    for ws in wb.worksheets:
        if ws.title in INVOICE_SHEETS:
            continue
        sheet = read_invoice_sheet(ws)
        key = accounts.find(ws.title)
        match = [i for i in invoices if i["account"] == key and i["invoice_date"] == sheet["date"]]
        if len(match) != 1:
            raise SystemExit(f"Onglet {ws.title} : facture du {sheet['date']} introuvable pour le compte {key}")
        invoice = match[0]
        day = "/".join(reversed(sheet["date"].split("-")))
        detail = [f"Facture{' n° ' + sheet['number'] if sheet['number'] else ''} du {day} (HT) :"]
        detail += [f"• {label} — {money(value)}" if value else f"{label} :" for label, value in sheet["lines"]]
        if sheet["discount"]:
            detail.append(f"Remise globale de {sheet['discount']:.0%}".replace("%", " %") + ".")
        if sheet["deposit"]:
            detail.append(f"Acompte de {money(sheet['deposit'][0])} reçu le {sheet['deposit'][1]}.")
        total = sum(v or 0 for _, v in sheet["lines"]) * (1 - (sheet["discount"] or 0))
        if abs(total - (invoice["amount"] or 0)) > 1:
            raise SystemExit(f"Onglet {ws.title} : total des lignes {total} ≠ montant facturé {invoice['amount']}")
        invoice["invoice_number"] = sheet["number"]
        invoice["comments"] = "\n\n".join(filter(None, [invoice.get("comments"), "\n".join(detail)]))
        client = sorted(sheet["client"], key=lambda v: v.startswith("À l'attention"))  # la raison sociale d'abord
        accounts.by_key[key]["billing_details"] = " · ".join(client)


def main(src, dest, invoices_src=None):
    wb = openpyxl.load_workbook(src, data_only=True)
    accounts = Accounts()

    invoices, opportunities, sheet_offers = [], [], {}
    for name in ACCOUNT_SHEETS:
        key, inv, offers, opps = read_account_sheet(wb[name], accounts)
        invoices += inv
        opportunities += opps
        sheet_offers[key] = offers

    ws = wb["Pipeline global"]
    deals = []
    for r in rows(ws, 9, 2, 22):
        if not isinstance(r[0], (int, float)):
            break
        kind = ACCOUNT_KINDS[text(r[1])]
        label = text(r[2])
        fiche = text(r[19])
        account = accounts.find(fiche.removeprefix("→").strip()) if fiche and fiche.strip() != "→ Appels d’offres" else None
        if account is None:
            account = accounts.add(re.sub(r"\s+", " ", label), kind)
        deals.append(
            {
                "number": int(r[0]),
                "account": account,
                "project_label": label,
                "title": text(r[3]),
                "solution": text(r[4]),
                "origin": text(r[5]),
                "lead": text(r[6]),
                "recurrence": text(r[7]),
                "stage": STAGES[text(r[8])],
                "closing": number(r[9]),
                "amount": number(r[10]),
                "commission": number(r[12]),
                "priority": priority(r[14]),
                "next_step": text(r[15]),
                "due_date": iso_date(r[16]),
                "comments": text(r[17]),
                "service": text(r[18]),
                "key_contact": text(r[20]),
            }
        )

    # Offres des fiches comptes absentes du pipeline : gardées comme opportunités du compte
    known = {(d["account"], d["title"]) for d in deals}
    for key, offers in sheet_offers.items():
        for offer in offers:
            titles = {t for (a, t) in known if a == key}
            if not any(offer["title"] in t or (t and t in offer["title"]) for t in titles) and offer["title"].startswith("Formation"):
                opportunities.append({"account": key, "title": offer["title"], "solution": offer["vendor"], "priority": None,
                                      "due_date": None, "owner": "Commercial IAT", "action": None, "comments": "Offre en préparation (fiche compte)."})

    ws = wb["Renouvellements"]
    renewals = []
    for r in rows(ws, 9, 2, 18):
        if not isinstance(r[0], (int, float)):
            break
        renewals.append(
            {
                "account": accounts.find(text(r[1])),
                "contract": text(r[2]),
                "last_invoice_date": iso_date(r[3]),
                "last_amount": number(r[4]),
                "history": text(r[5]),
                "due_date_override": None,
                "status": text(r[11]) or "Non démarré",
                "owner": text(r[12]),
                "last_reminder_date": iso_date(r[13]),
                "next_action": text(r[14]),
                "comments": text(r[15]),
            }
        )

    ws = wb["Appels d’offres"]
    tenders = []
    for r in rows(ws, 9, 2, 15):
        if not isinstance(r[0], (int, float)):
            break
        tenders.append(
            {
                "organization": text(r[1]),
                "subject": text(r[2]),
                "reference": text(r[3]),
                "published_on": iso_date(r[4]),
                "deadline": iso_date(r[5]),
                "estimated_amount": number(r[7]),
                "response_mode": text(r[8]),
                "partners": text(r[9]),
                "requirements": text(r[10]),
                "stage": text(r[11]),
                "owner": text(r[12]),
                "comments": text(r[13]),
            }
        )

    ws = wb["Solutions & prospects"]
    solutions = []
    for r in rows(ws, 9, 2, 13):
        solutions.append(
            {
                "name": text(r[0]),
                "vendor": text(r[1]),
                "category": text(r[2]),
                "partnership_status": text(r[3]),
                "internal_expert": text(r[4]),
                "mastery": text(r[5]),
                "target_sectors": text(r[6]),
                "value_pitch": text(r[7]),
                "existing_clients": text(r[8]),
                "materials": text(r[9]),
                "priority": priority(r[10]),
                "next_action": text(r[11]),
            }
        )
    prospects = []
    start = find_row(ws, lambda v: v.startswith("2. PROSPECTS"))
    for r in rows(ws, start + 2, 2, 13):
        if not text(r[1]):
            continue
        prospects.append(
            {
                "solution": text(r[0]),
                "company": text(r[1]),
                "sector": text(r[2]),
                "contact": text(r[3]),
                "angle": text(r[4]),
                "stage": STAGES.get(text(r[5])),
                "attack_week": iso_date(r[6]),
                "next_action": text(r[7]),
                "owner": text(r[8]),
                "relationship_score": number(r[9]),
                "last_action_date": iso_date(r[10]),
                "comments": text(r[11]),
            }
        )
    pitches = []
    start = find_row(ws, lambda v: v.startswith("3. PITCHS"))
    for r in rows(ws, start + 2, 2, 12):
        pitches.append(
            {
                "solution": "Credolab",
                "segment": text(r[0]),
                "promise": text(r[1]),
                "stake": text(r[2]),
                "offer": text(r[4]),
                "hook": text(r[7]),
                "call_to_action": text(r[8]),
                "targets": text(r[10]),
            }
        )

    ws = wb["Plan d’actions"]
    actions = []
    for r in rows(ws, 9, 2, 9):
        if not isinstance(r[0], (int, float)):
            break
        actions.append(
            {
                "number": int(r[0]),
                "title": text(r[1]),
                "related_to": text(r[2]),
                "account": accounts.find(text(r[2])) if text(r[2]) and len(text(r[2])) < 40 else None,
                "owner": text(r[3]),
                "due_date": iso_date(r[4]),
                "priority": priority(r[5]),
                "status": text(r[6]) or "À faire",
                "comments": text(r[7]),
            }
        )

    ws = wb["Contacts"]
    contacts = []
    for r in rows(ws, 9, 2, 14):
        if not isinstance(r[0], (int, float)):
            break
        link = text(r[8])
        is_research = (text(r[12]) or "").startswith("Recherche")
        contacts.append(
            {
                "company": text(r[1]),
                "sector": text(r[2]),
                "name": text(r[3]),
                "role": text(r[4]),
                "profile": None if text(r[5]) == "Non renseigné" else text(r[5]),
                "phone": text(r[6]),
                "email": text(r[7]),
                "pipeline_link": link,
                "account": accounts.find(link) or (accounts.find(text(r[1])) if is_research else None),
                "status": text(r[9]),
                "last_contact": iso_date(r[10]),
                "comments": text(r[11]),
                "source": "research" if is_research else "client_base",
            }
        )

    ws = wb["Listes"]
    lists = {}
    for col in ws.iter_cols(min_row=1):
        values = [text(c.value) for c in col if text(c.value)]
        if values:
            lists[values[0]] = values[1:]

    if invoices_src:
        enrich_from_invoices(invoices_src, accounts, invoices)

    data = {
        "meta": {"source": " + ".join(p.split("/")[-1] for p in filter(None, [src, invoices_src])), "imported_at": datetime.now().isoformat(timespec="seconds")},
        "accounts": list(accounts.by_key.values()),
        "deals": deals,
        "invoices": invoices,
        "renewals": renewals,
        "opportunities": opportunities,
        "contacts": contacts,
        "tenders": tenders,
        "solutions": solutions,
        "prospects": prospects,
        "pitches": pitches,
        "actions": actions,
        "lists": lists,
    }
    with open(dest, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=1)
    print({k: len(v) for k, v in data.items() if isinstance(v, list)})


if __name__ == "__main__":
    main(*sys.argv[1:4])
